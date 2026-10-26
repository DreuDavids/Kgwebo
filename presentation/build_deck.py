"""Build the Kgwebo FNB App of the Year 2026 16:9 deck from real screenshots."""

from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
from lxml import etree

ROOT = Path(__file__).resolve().parent
REPO = ROOT.parent
SHOTS = ROOT / "assets" / "screenshots"
FRAMED = ROOT / "assets" / "framed"
ICON = REPO / "brand" / "kit" / "icon-512.png"
OUT_PPTX = ROOT / "Kgwebo_FNB_App_of_the_Year_2026.pptx"

INK = RGBColor(0x11, 0x11, 0x11)
LIME = RGBColor(0xC6, 0xFF, 0x4A)
CANVAS = RGBColor(0xF4, 0xF5, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0x6B, 0x72, 0x80)
FAINT = RGBColor(0x9C, 0xA3, 0xAF)
LINE = RGBColor(0xE6, 0xE8, 0xEC)
AMBER = RGBColor(0xC4, 0x84, 0x1D)
SANS = "Plus Jakarta Sans"
SERIF = "Fraunces"

W = Inches(13.333)
H = Inches(7.5)


def _set_run_font(run, name: str, size: int, bold: bool, color: RGBColor) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        el = rPr.find(qn(f"a:{tag}"))
        if el is None:
            el = etree.SubElement(rPr, qn(f"a:{tag}"))
        el.set("typeface", name)


def add_text(
    slide,
    left,
    top,
    width,
    height,
    text: str,
    *,
    size: int = 16,
    bold: bool = False,
    color: RGBColor = INK,
    font: str = SANS,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
) -> None:
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    try:
        tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    _set_run_font(run, font, size, bold, color)


def add_rect(slide, left, top, width, height, fill: RGBColor, radius: int | None = None):
    shape = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shape, left, top, width, height)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    if radius:
        s.adjustments[0] = 0.12
    return s


def notes(slide, text: str) -> None:
    slide.notes_slide.notes_text_frame.text = text


def frame_phone(src: Path, dest: Path, radius: int = 72) -> Path:
    dest.parent.mkdir(parents=True, exist_ok=True)
    im = Image.open(src).convert("RGBA")
    mask = Image.new("L", im.size, 0)
    ImageDraw.Draw(mask).rounded_rectangle((0, 0) + im.size, radius=radius, fill=255)
    im.putalpha(mask)
    bg = Image.new("RGBA", im.size, (17, 17, 17, 255))
    ImageDraw.Draw(bg).rounded_rectangle((0, 0) + im.size, radius=radius, fill=(17, 17, 17, 255))
    composed = Image.alpha_composite(bg, im)
    composed.save(dest)
    return dest


def phone(slide, src: Path, left, top, height) -> None:
    framed = FRAMED / src.name
    if not framed.exists() or framed.stat().st_mtime < src.stat().st_mtime:
        frame_phone(src, framed)
    im = Image.open(framed)
    ratio = im.width / im.height
    width = int(height * ratio)
    slide.shapes.add_picture(str(framed), left, top, width=width, height=height)


def footer(slide, page: int) -> None:
    add_rect(slide, Inches(0), H - Inches(0.38), W, Inches(0.38), WHITE)
    add_text(
        slide,
        Inches(0.5),
        H - Inches(0.34),
        Inches(8),
        Inches(0.28),
        "Kgwebo  ·  FNB App of the Year 2026  ·  Judge demo: dreu.davids1@gmail.com",
        size=10,
        color=MUTED,
    )
    add_text(
        slide,
        W - Inches(1.2),
        H - Inches(0.34),
        Inches(0.7),
        Inches(0.28),
        str(page),
        size=10,
        color=FAINT,
        align=PP_ALIGN.RIGHT,
    )


def chip(slide, left, top, width, height, label: str, value: str, fill: RGBColor = WHITE) -> None:
    box = add_rect(slide, left, top, width, height, fill, radius=1)
    box.adjustments[0] = 0.1
    add_text(slide, left + Inches(0.16), top + Inches(0.1), width - Inches(0.3), Inches(0.24), label, size=11, color=MUTED)
    add_text(slide, left + Inches(0.16), top + Inches(0.34), width - Inches(0.3), Inches(0.4), value, size=16, bold=True, color=INK)


def badge(slide, left, top, number: str) -> None:
    add_rect(slide, left, top, Inches(0.34), Inches(0.34), INK, radius=1)
    add_text(slide, left, top + Inches(0.02), Inches(0.34), Inches(0.3), number, size=12, bold=True, color=LIME, align=PP_ALIGN.CENTER)


def down_mark(slide, left, top) -> None:
    add_text(slide, left, top, Inches(1.2), Inches(0.28), "↓", size=16, bold=True, color=INK, align=PP_ALIGN.CENTER)


def creds_box(slide, left, top, width, height, large: bool = False) -> None:
    box = add_rect(slide, left, top, width, height, INK, radius=1)
    box.adjustments[0] = 0.08
    title_size = 13 if large else 11
    body_size = 16 if large else 13
    add_text(slide, left + Inches(0.22), top + Inches(0.12), width - Inches(0.4), Inches(0.28), "JUDGE DEMO ACCOUNT", size=title_size, bold=True, color=LIME)
    add_text(
        slide,
        left + Inches(0.22),
        top + Inches(0.38),
        width - Inches(0.4),
        height - Inches(0.5),
        "Email: dreu.davids1@gmail.com\nPassword: password\nStall: DreuTech  ·  illustrative records",
        size=body_size,
        color=WHITE,
    )


def new_slide(prs, fill: RGBColor = CANVAS):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(slide, 0, 0, W, H, fill)
    return slide


def build() -> Path:
    FRAMED.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H

    dash = SHOTS / "01-dashboard.png"
    sale = SHOTS / "02-record-sale.png"
    money = SHOTS / "03-money-history.png"
    bills = SHOTS / "04-cash-flow.png"
    stock = SHOTS / "05-stock.png"
    recs = SHOTS / "06-recommendations.png"
    ask_intro = SHOTS / "07-ask-intro.png"
    ask_cash = SHOTS / "07-ask-cash.png"
    ask_draft = SHOTS / "08-ask-draft.png"
    reports = SHOTS / "09-reports.png"
    welcome = SHOTS / "00-welcome.png"

    # 1 Cover
    s = new_slide(prs, INK)
    s.shapes.add_picture(str(ICON), Inches(0.55), Inches(0.45), Inches(0.62), Inches(0.62))
    add_text(s, Inches(1.3), Inches(0.5), Inches(5), Inches(0.5), "Kgwebo", size=28, bold=True, color=LIME, font=SERIF)
    add_text(s, Inches(0.55), Inches(1.45), Inches(7.2), Inches(1.3), "Helping small businesses\nunderstand their money.", size=36, bold=True, color=WHITE)
    add_text(
        s,
        Inches(0.55),
        Inches(3.0),
        Inches(6.8),
        Inches(1.1),
        "A Flutter stall-finance app that records cash-basis activity, shows what cash is actually available, and turns those records into the next useful step.",
        size=16,
        color=RGBColor(0xD1, 0xD5, 0xDB),
    )
    add_text(s, Inches(0.55), Inches(4.25), Inches(6.5), Inches(0.4), "2026 FNB App of the Year  ·  Youth entrepreneurship & small-business growth", size=13, color=LIME)
    creds_box(s, Inches(0.55), Inches(4.9), Inches(6.4), Inches(1.55), large=True)
    if dash.exists():
        phone(s, dash, Inches(8.35), Inches(0.35), Inches(6.8))
    notes(
        s,
        "Open on the product, not awards. Credentials are required for judges. DreuTech is an illustrative stall in ZAR. Do not claim traction.",
    )

    # 2 Problem
    s = new_slide(prs)
    add_text(s, Inches(0.55), Inches(0.35), Inches(12), Inches(0.35), "THE PROBLEM", size=12, bold=True, color=MUTED)
    add_text(s, Inches(0.55), Inches(0.7), Inches(12), Inches(1.0), "The till is busy. The picture of the money is not.", size=32, bold=True, color=INK)
    add_text(
        s,
        Inches(0.55),
        Inches(1.7),
        Inches(12),
        Inches(0.7),
        "These are practical problems for informal traders and small service owners across Southern Africa. They are product hypotheses, not a market-size claim.",
        size=15,
        color=MUTED,
    )
    cards = [
        ("Records live in notebooks\nand chat threads", "Sales, expenses and stock are captured inconsistently, if at all."),
        ("Cash on hand is a guess", "Money in the till is mixed with personal cash, unpaid sales and unpaid bills."),
        ("Upcoming bills arrive as surprises", "Rent, software and supplier payments are not visible before they fall due."),
        ("Numbers rarely become a decision", "A tracker that only lists transactions does not say what to do next."),
    ]
    for i, (title, body) in enumerate(cards):
        x = Inches(0.55) + Inches(i * 3.15)
        add_rect(s, x, Inches(2.7), Inches(2.95), Inches(3.7), WHITE, radius=1)
        add_rect(s, x, Inches(2.7), Inches(0.14), Inches(3.7), LIME)
        add_text(s, x + Inches(0.28), Inches(2.95), Inches(2.5), Inches(1.4), title, size=16, bold=True, color=INK)
        add_text(s, x + Inches(0.28), Inches(4.5), Inches(2.5), Inches(1.6), body, size=13, color=MUTED)
    footer(s, 2)
    notes(s, "Stay human. No invented statistics. Botswana-first product; screenshots later show a ZA/ZAR illustrative stall.")

    # 3 Solution
    s = new_slide(prs)
    add_text(s, Inches(0.55), Inches(0.35), Inches(7), Inches(0.3), "THE SOLUTION", size=12, bold=True, color=MUTED)
    add_text(s, Inches(0.55), Inches(0.65), Inches(7.2), Inches(1.1), "One place to record, understand, anticipate and act.", size=28, bold=True, color=INK)
    add_text(
        s,
        Inches(0.55),
        Inches(1.85),
        Inches(7.1),
        Inches(0.9),
        "Kgwebo is a working Android app. It already records sales and expenses, keeps a stock catalogue, forecasts outgoing bills, and answers questions from recorded figures — not from invented totals.",
        size=15,
        color=MUTED,
    )
    points = [
        ("Record", "Paid and unpaid sales, expenses, refunds, owner cash moves."),
        ("Understand", "Recorded cash is separated from sales. Opening cash is not revenue."),
        ("Anticipate", "7 / 14 / 30-day outlook subtracts unpaid bills. Incoming cash is not invented."),
        ("Act", "Deterministic recommendations plus Ask, which drafts but never auto-saves."),
    ]
    for i, (title, body) in enumerate(points):
        y = Inches(2.9) + Inches(i * 0.95)
        add_rect(s, Inches(0.55), y, Inches(7.1), Inches(0.85), WHITE, radius=1)
        add_text(s, Inches(0.75), y + Inches(0.1), Inches(6.7), Inches(0.3), title, size=15, bold=True, color=INK)
        add_text(s, Inches(0.75), y + Inches(0.4), Inches(6.7), Inches(0.4), body, size=13, color=MUTED)
    if dash.exists():
        phone(s, dash, Inches(8.45), Inches(0.55), Inches(6.2))
    footer(s, 3)
    notes(s, "State clearly what works: live auth, dashboard, money, stock, bills, Ask, reports. Demo stall is DreuTech in ZAR.")

    # 4 How it works
    s = new_slide(prs)
    add_text(s, Inches(0.55), Inches(0.28), Inches(12), Inches(0.28), "HOW IT WORKS", size=12, bold=True, color=MUTED)
    add_text(s, Inches(0.55), Inches(0.55), Inches(12), Inches(0.55), "Four screens. Four jobs.", size=28, bold=True, color=INK)
    shots = [
        (money, "Money", "Sales, expenses, refunds and owner movements in one ledger."),
        (stock, "Stock", "Catalogue, cost, quantity. Inventory value is not cash."),
        (bills, "Bills", "Unpaid commitments sit in the forecast, not in recorded cash."),
        (ask_intro, "Ask", "Talk or type. Figures come from Kgwebo. Nothing saves until confirm."),
    ]
    for i, (path, title, caption) in enumerate(shots):
        x = Inches(0.4) + Inches(i * 3.2)
        if path.exists():
            phone(s, path, x + Inches(0.35), Inches(1.2), Inches(4.15))
        add_text(s, x, Inches(5.45), Inches(3.0), Inches(0.32), title, size=16, bold=True, color=INK, align=PP_ALIGN.CENTER)
        add_text(s, x, Inches(5.78), Inches(3.0), Inches(0.85), caption, size=12, color=MUTED, align=PP_ALIGN.CENTER)
    footer(s, 4)
    notes(
        s,
        "Offline: only amount-only creates and new unpaid bills queue. Pending is not cash. Product sales need a network. Ask can fall back to on-device heuristics if the cloud call fails.",
    )

    # 5 Journey — the incident
    s = new_slide(prs)
    add_text(s, Inches(0.55), Inches(0.28), Inches(8), Inches(0.28), "USER STORY  ·  ILLUSTRATIVE DREUTECH", size=12, bold=True, color=MUTED)
    add_text(s, Inches(0.55), Inches(0.55), Inches(8.0), Inches(1.15), "The till looks full.\nRent is due in six days.", size=30, bold=True, color=INK)
    add_text(
        s,
        Inches(0.55),
        Inches(1.8),
        Inches(7.9),
        Inches(1.55),
        "Monday 26 October. Dreu runs DreuTech — cables, laptops, and the odd website job. A supplier WhatsApps: ten USB-C hubs, R 4,000 cash, today. The till shows R 31,840. Instinct says yes.\n\nWhat the till does not say: office rent of R 5,500 lands on 1 November. Adobe Creative Cloud of R 890 lands on 5 November. Three products are already low. A client just paid R 450 for a website fix that is not in the book yet.",
        size=14,
        color=MUTED,
    )
    metrics = [
        ("Recorded cash", "R 31,840"),
        ("Bills in 14 days", "R 6,390"),
        ("Supplier offer", "R 4,000"),
        ("Cash after bills", "R 25,450"),
    ]
    for i, (label, value) in enumerate(metrics):
        col, row = i % 2, i // 2
        x = Inches(0.55) + Inches(col * 3.95)
        y = Inches(3.55) + Inches(row * 1.15)
        chip(s, x, y, Inches(3.75), Inches(1.02), label, value)
    add_text(
        s,
        Inches(0.55),
        Inches(5.95),
        Inches(7.9),
        Inches(0.85),
        "The question is not “is there money in the till?” It is “can Dreu restock and still pay rent on 1 November?” Illustrative stall — not a case study, not guaranteed savings.",
        size=13,
        color=INK,
    )
    if bills.exists():
        phone(s, bills, Inches(8.85), Inches(0.4), Inches(6.55))
    footer(s, 5)
    notes(
        s,
        "Tell the WhatsApp restock offer first. Point at the bills drop: R 31,840 → R 25,450. Label DreuTech illustrative. Do not claim Dreu is a real customer.",
    )

    # 6 Journey — how they used the app
    s = new_slide(prs)
    add_text(s, Inches(0.5), Inches(0.22), Inches(12), Inches(0.24), "USER STORY  ·  SAME MORNING", size=12, bold=True, color=MUTED)
    add_text(s, Inches(0.5), Inches(0.46), Inches(12.2), Inches(0.45), "Dreu opened Kgwebo and answered it.", size=24, bold=True, color=INK)
    stages = [
        ("See the cash", "Home shows R 31,840 recorded cash. Unpaid bills are excluded. This is not free-to-spend money.", dash),
        ("Subtract the bills", "14-day Bills path: rent 1 Nov, Adobe 5 Nov. After known bills: R 25,450. Incoming cash is not invented.", bills),
        ("Check the shelf", "USB-C hub is genuinely low (2 left). Stock on the shelf is R 69,490 — labelled as not cash.", stock),
        ("Ask, then record", "Ask answers cash from records. “I received R 450 for a website fix” becomes a draft. Confirm still waiting.", ask_draft if ask_draft.exists() else ask_intro),
    ]
    for i, (title, body, img) in enumerate(stages):
        x = Inches(0.35) + Inches(i * 3.22)
        add_rect(s, x, Inches(1.02), Inches(3.08), Inches(5.5), WHITE, radius=1)
        badge(s, x + Inches(0.16), Inches(1.16), str(i + 1))
        add_text(s, x + Inches(0.56), Inches(1.16), Inches(2.35), Inches(0.34), title, size=13, bold=True, color=INK)
        add_text(s, x + Inches(0.16), Inches(1.54), Inches(2.76), Inches(1.15), body, size=11, color=MUTED)
        if img.exists():
            phone(s, img, x + Inches(0.52), Inches(2.72), Inches(3.55))
    add_text(
        s,
        Inches(0.5),
        Inches(6.58),
        Inches(12.2),
        Inches(0.42),
        "Decision: R 4,000 restock still fits inside R 25,450 after rent and Adobe. Dreu can say yes with eyes open. The R 450 sale is not in cash until Confirm.",
        size=13,
        color=INK,
    )
    footer(s, 6)
    notes(
        s,
        "Walk the four screens as one morning. End on the decision: restock is affordable after bills. Ask’s Confirm is still unsaved — say that out loud.",
    )

    # 7 Product close-up
    s = new_slide(prs)
    add_text(s, Inches(0.45), Inches(0.22), Inches(12), Inches(0.28), "A CLOSER LOOK", size=12, bold=True, color=MUTED)
    add_text(s, Inches(0.45), Inches(0.48), Inches(12), Inches(0.45), "The product, as it runs.", size=26, bold=True, color=INK)
    grid = [
        (dash, "Home  ·  recorded cash first"),
        (money, "Money  ·  the actual ledger"),
        (bills, "Bills  ·  cash after known commitments"),
        (ask_draft, "Ask  ·  draft, then confirm"),
    ]
    for i, (img, label) in enumerate(grid):
        x = Inches(0.25) + Inches(i * 3.25)
        y = Inches(1.05)
        if img.exists():
            phone(s, img, x + Inches(0.28), y, Inches(5.15))
        add_text(s, x, Inches(6.35), Inches(3.2), Inches(0.4), label, size=12, bold=True, color=INK, align=PP_ALIGN.CENTER)
    footer(s, 7)
    notes(s, "Let the screenshots do the talking. Point at cash vs sales, the forecast drop, and Ask’s Confirm button.")

    # 8 Intelligence
    s = new_slide(prs)
    add_text(s, Inches(0.5), Inches(0.28), Inches(12), Inches(0.25), "INTELLIGENCE", size=12, bold=True, color=MUTED)
    add_text(s, Inches(0.5), Inches(0.52), Inches(7.4), Inches(0.9), "Trustworthy figures first. Then a copilot that cannot quietly change them.", size=24, bold=True, color=INK)
    facts = [
        ("Cash is not sales", "Recorded cash = opening cash + paid in − paid out. Dreu’s R 24,000 laptop sale is cash only because it is marked paid."),
        ("The till is not free cash", "R 31,840 is recorded cash. R 6,390 of bills still sit outside it until they are paid."),
        ("Bills wait their turn", "Unpaid rent does not reduce cash today. The 14-day outlook subtracts it. Incoming cash is never invented."),
        ("Recommendations are rules", "Six evidence-backed rules. Low stock uses catalogue quantities. Copy says estimate, not advice."),
        ("Ask is grounded", "Ask reads FinancialEngine. Money talk is overwritten with recorded figures. Nothing saves until Confirm."),
        ("Stock value is not cash", "R 69,490 on the shelf is cost × quantity. Spending it would be spending goods Dreu still needs to sell."),
    ]
    for i, (title, body) in enumerate(facts):
        col, row = i % 2, i // 2
        x = Inches(0.5) + Inches(col * 4.15)
        y = Inches(1.55) + Inches(row * 1.55)
        add_rect(s, x, y, Inches(3.95), Inches(1.4), WHITE, radius=1)
        add_text(s, x + Inches(0.2), y + Inches(0.12), Inches(3.55), Inches(0.35), title, size=14, bold=True, color=INK)
        add_text(s, x + Inches(0.2), y + Inches(0.5), Inches(3.55), Inches(0.78), body, size=12, color=MUTED)
    if ask_draft.exists():
        phone(s, ask_draft, Inches(9.05), Inches(1.45), Inches(5.2))
    footer(s, 8)
    notes(
        s,
        "Tie numbers back to Dreu’s morning. Do not imply a silent LLM posting. Cloud Ask uses kgwebo-copilot (LiteLLM). If the key is missing, Ask still answers from recorded totals.",
    )

    # 9 Architecture — the stack
    s = new_slide(prs)
    add_text(s, Inches(0.5), Inches(0.2), Inches(12), Inches(0.24), "ARCHITECTURE", size=12, bold=True, color=MUTED)
    add_text(s, Inches(0.5), Inches(0.42), Inches(12.2), Inches(0.5), "Built so the phone cannot lie about money.", size=26, bold=True, color=INK)

    def arch_cell(x, y, w, h, title, body, fill=WHITE, title_color=INK, body_color=MUTED):
        box = add_rect(s, x, y, w, h, fill, radius=1)
        box.adjustments[0] = 0.08
        add_text(s, x + Inches(0.12), y + Inches(0.08), w - Inches(0.22), Inches(0.28), title, size=12, bold=True, color=title_color)
        add_text(s, x + Inches(0.12), y + Inches(0.36), w - Inches(0.22), h - Inches(0.44), body, size=11, color=body_color)

    def row_of(y, items, fill=WHITE, title_color=INK, body_color=MUTED):
        n = len(items)
        gap = 0.12
        total = 10.55
        w = (total - gap * (n - 1)) / n
        for i, (title, body) in enumerate(items):
            x = Inches(2.25) + Inches(i * (w + gap))
            arch_cell(x, y, Inches(w), Inches(1.02), title, body, fill=fill, title_color=title_color, body_color=body_color)

    add_text(s, Inches(0.45), Inches(1.12), Inches(1.7), Inches(0.9), "PRESENTATION", size=10, bold=True, color=MUTED)
    row_of(
        Inches(1.05),
        [
            ("Flutter UI", "Android screens. Widgets render values. They never add cash."),
            ("GoRouter", "Auth guards and onboarding redirects. Hiding a button is not the lock."),
            ("Riverpod", "One graph of state. Home, Bills, Reports and Ask subscribe to the same summary."),
        ],
    )
    down_mark(s, Inches(6.9), Inches(2.05))
    add_text(s, Inches(0.45), Inches(2.42), Inches(1.7), Inches(0.9), "MONEY CORE", size=10, bold=True, color=MUTED)
    row_of(
        Inches(2.32),
        [
            ("FinancialEngine", "Pure Dart. Integer minor units. Opening cash is not revenue."),
            ("CashForecast", "7 / 14 / 30-day path. Subtracts unpaid bills. Incoming cash is never invented."),
            ("RecommendationRules", "Six deterministic rules. Empty stalls stay quiet."),
            ("CopilotOrchestrator", "Classifies intent, grounds replies, never writes the ledger."),
        ],
    )
    down_mark(s, Inches(6.9), Inches(3.32))
    add_text(s, Inches(0.45), Inches(3.68), Inches(1.7), Inches(0.9), "ON THE PHONE", size=10, bold=True, color=MUTED)
    row_of(
        Inches(3.58),
        [
            ("Drift cache", "Owner-keyed confirmed rows only. Stale cache is labelled, never invented."),
            ("Sync outbox", "Amount-only creates and new bills. Pending is not cash, sales, or a forecast input."),
        ],
    )
    down_mark(s, Inches(6.9), Inches(4.58))
    add_text(s, Inches(0.45), Inches(4.95), Inches(1.7), Inches(0.9), "SUPABASE", size=10, bold=True, color=MUTED)
    row_of(
        Inches(4.85),
        [
            ("Auth", "Email/password. JWT on every call. Publishable key only in the APK."),
            ("Postgres + FORCE RLS", "Ownership frozen on write. A second account cannot read DreuTech."),
            ("Inventory RPCs", "record_product_sale is atomic. Clients cannot insert stock rows."),
            ("kgwebo-copilot", "Edge Function. Rate-limited. LiteLLM JSON. No SQL from the model."),
        ],
        fill=INK,
        title_color=LIME,
        body_color=RGBColor(0xD1, 0xD5, 0xDB),
    )
    add_text(
        s,
        Inches(0.5),
        Inches(6.05),
        Inches(12.3),
        Inches(0.85),
        "One engine feeds Home, Bills, Reports, the PDF and Ask. Isolation lives in Postgres (FORCE RLS, frozen owner_id, SECURITY DEFINER RPCs) — not in the client. There are no bank feeds, payment processors, or silent auto-posts in this build.",
        size=13,
        color=INK,
    )
    footer(s, 9)
    notes(
        s,
        "Walk the stack top to bottom. Emphasise: widgets do not calculate; pending is not cash; RLS not client honour; Ask cannot write SQL. Do not draw bank feeds.",
    )

    # 10 Architecture — two pipelines
    s = new_slide(prs)
    add_text(s, Inches(0.5), Inches(0.2), Inches(12), Inches(0.24), "ARCHITECTURE  ·  TWO PATHS", size=12, bold=True, color=MUTED)
    add_text(s, Inches(0.5), Inches(0.42), Inches(12.2), Inches(0.5), "The model is a parser. The ledger is the source of truth.", size=24, bold=True, color=INK)

    add_rect(s, Inches(0.45), Inches(1.05), Inches(6.05), Inches(5.45), WHITE, radius=1)
    add_rect(s, Inches(0.45), Inches(1.05), Inches(0.14), Inches(5.45), LIME)
    add_text(s, Inches(0.8), Inches(1.18), Inches(5.5), Inches(0.35), "When Dreu records a sale", size=16, bold=True, color=INK)
    record_steps = [
        ("1", "Form validators accept integer minor units. Binary float never holds money."),
        ("2", "Repository tries Supabase first. Auth JWT. Idempotency key on every create."),
        ("3", "Amount-only sale or new bill, network fails → durable outbox. Status: pending. FinancialEngine ignores it."),
        ("4", "Product sale stays online. record_product_sale writes the transaction, line and stock movement in one RPC. Retry cannot double-count."),
        ("5", "Confirmed row upserts Drift. Home, Bills, Reports, PDF and Ask all recompute from FinancialEngine."),
    ]
    for i, (num, body) in enumerate(record_steps):
        y = Inches(1.62) + Inches(i * 0.9)
        badge(s, Inches(0.8), y, num)
        add_text(s, Inches(1.28), y - Inches(0.04), Inches(4.95), Inches(0.85), body, size=13, color=MUTED)

    add_rect(s, Inches(6.75), Inches(1.05), Inches(6.05), Inches(5.45), INK, radius=1)
    add_text(s, Inches(7.1), Inches(1.18), Inches(5.5), Inches(0.35), "When Dreu asks a question", size=16, bold=True, color=LIME)
    ask_steps = [
        ("1", "CopilotOrchestrator classifies intent from the utterance."),
        ("2", "A snapshot is read from FinancialEngine — the same R 31,840 Home already shows."),
        ("3", "JWT call to kgwebo-copilot. Rate-limited Edge Function. LiteLLM returns JSON only."),
        ("4", "groundCopilotTurn overwrites any invented cash, sales or advice with recorded figures."),
        ("5", "A draft is shown. Confirm runs CopilotLedgerCommit through the same repository as a typed sale. Cloud down? On-device heuristics, still grounded, still confirm-before-save."),
    ]
    for i, (num, body) in enumerate(ask_steps):
        y = Inches(1.62) + Inches(i * 0.9)
        add_rect(s, Inches(7.1), y, Inches(0.34), Inches(0.34), LIME, radius=1)
        add_text(s, Inches(7.1), y + Inches(0.02), Inches(0.34), Inches(0.3), num, size=12, bold=True, color=INK, align=PP_ALIGN.CENTER)
        add_text(s, Inches(7.58), y - Inches(0.04), Inches(4.95), Inches(0.85), body, size=13, color=RGBColor(0xE5, 0xE7, 0xEB))
    footer(s, 10)
    notes(
        s,
        "Left path: honest offline. Right path: LLM cannot mutate the ledger. Say ‘confirm-before-save’ twice. Do not claim Ask is always online.",
    )

    # 11 Impact
    s = new_slide(prs)
    add_text(s, Inches(0.55), Inches(0.32), Inches(12), Inches(0.28), "IMPACT AND SUSTAINABILITY", size=12, bold=True, color=MUTED)
    add_text(s, Inches(0.55), Inches(0.65), Inches(12), Inches(0.8), "Better habits. Clearer cash. Room to grow.", size=28, bold=True, color=INK)
    impacts = [
        ("Daily recording", "A stall that captures the R 450 website fix the same morning can see cash move instead of reconstructing the week from chat threads."),
        ("Cash that is actually available", "Dreu’s restock decision uses R 25,450 after known bills, not the R 31,840 sitting in the till. Money already spoken for stays visible."),
        ("Stock discipline", "Low-stock notices use saved catalogue quantities. Inventory value is labelled as not cash, so a full shelf does not look like a full till."),
        ("A path that stays cheap to run", "One Flutter app, one Supabase project, one money engine. Operating cost stays a small backend, not a stack of bank feeds and extra models."),
    ]
    for i, (title, body) in enumerate(impacts):
        y = Inches(1.65) + Inches(i * 1.2)
        add_rect(s, Inches(0.55), y, Inches(12.2), Inches(1.08), WHITE, radius=1)
        add_text(s, Inches(0.8), y + Inches(0.16), Inches(11.7), Inches(0.3), title, size=16, bold=True, color=INK)
        add_text(s, Inches(0.8), y + Inches(0.5), Inches(11.7), Inches(0.48), body, size=14, color=MUTED)
    footer(s, 11)
    notes(s, "No testimonials, no user counts, no revenue projections. Tie impact back to Dreu’s morning. Potential outcomes only.")

    # 12 Close
    s = new_slide(prs, INK)
    s.shapes.add_picture(str(ICON), Inches(0.6), Inches(0.45), Inches(0.7), Inches(0.7))
    add_text(s, Inches(1.45), Inches(0.55), Inches(5), Inches(0.5), "Kgwebo", size=28, bold=True, color=LIME, font=SERIF)
    add_text(s, Inches(0.6), Inches(1.5), Inches(7.4), Inches(1.6), "Every business deserves a clearer picture of its money.", size=34, bold=True, color=WHITE)
    add_text(
        s,
        Inches(0.6),
        Inches(3.3),
        Inches(7.2),
        Inches(1.0),
        "Record what happened. See the cash that is actually free. Act before rent lands.",
        size=16,
        color=RGBColor(0xD1, 0xD5, 0xDB),
    )
    creds_box(s, Inches(0.6), Inches(4.55), Inches(7.0), Inches(1.85), large=True)
    hero = ask_cash if ask_cash.exists() else dash
    if hero.exists():
        phone(s, hero, Inches(8.4), Inches(0.4), Inches(6.7))
    notes(s, "End on the line. Point judges at the credentials. Offer to open Ask live if time remains.")

    prs.save(OUT_PPTX)
    return OUT_PPTX


if __name__ == "__main__":
    path = build()
    print(f"wrote {path}")
